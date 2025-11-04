"""
Fixtures para criar arquivos Excel de teste
"""
from django.core.files.uploadedfile import SimpleUploadedFile
import openpyxl
import io


class ExcelFixtures:
    """Classe para gerar arquivos Excel de teste"""

    @staticmethod
    def create_simple_excel():
        """Cria um arquivo Excel simples para testes básicos"""
        wb = openpyxl.Workbook()
        ws = wb.active

        # Headers
        ws['A1'] = 'Código'
        ws['B1'] = 'Descrição'
        ws['C1'] = 'Preço'

        # Dados
        ws['A2'] = 'P001'
        ws['B2'] = 'Produto 1'
        ws['C2'] = 100.00

        ws['A3'] = 'P002'
        ws['B3'] = 'Produto 2'
        ws['C3'] = 200.00

        # Salvar em BytesIO
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)

        return SimpleUploadedFile(
            "test_simple.xlsx",
            excel_buffer.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    @staticmethod
    def create_complex_excel():
        """
        Cria arquivo Excel complexo similar à planilha Dona Flor
        """
        wb = openpyxl.Workbook()
        ws = wb.active

        # Headers na linha 12 (index 11)
        headers = [
            'Foto', 'Num', 'Código', 'Descrição', 'Dimensões',
            'Cúbico', 'Peso', 'NCM', 'Preço 1', 'Preço 2', 'Preço 3'
        ]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=12, column=col_idx, value=header)

        # Dados - Produto com variantes
        data_rows = [
            # Produto 1: Almofada com 3 variantes
            [None, 1, 'AC2', 'Almofada Decorativa', 'A15 * L45 * P45',
             0.037553, 0.55, '94049000', 123.12, 136.0, 150.55],
            [None, 2, 'AC3', None, 'A15 * L55 * P55',
             0.055233, 0.85, '94049000', 157.15, 173.0, 191.01],
            [None, 3, 'AC33', None, 'A15 * L60 * P60',
             0.066, 1.20, '94049000', 189.00, 205.0, 225.50],

            # Produto 2: Banqueta com 2 variantes
            [None, 4, 'BC1', 'Banqueta Alta', 'A30 * L30 * P30',
             0.027, 2.50, '94036000', 250.00, 275.0, 302.50],
            [None, 5, 'BC2', None, 'A35 * L35 * P35',
             0.042, 3.20, '94036000', 310.00, 341.0, 375.10],

            # Produto 3: Mesa sem variantes
            [None, 6, 'M1', 'Mesa de Centro', 'A45 * L120 * P60',
             0.324, 15.00, '94036000', 850.00, 935.0, 1028.50],
        ]

        for row_idx, row_data in enumerate(data_rows, start=13):
            for col_idx, value in enumerate(row_data, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        # Salvar em BytesIO
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)

        return SimpleUploadedFile(
            "test_complex.xlsx",
            excel_buffer.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    @staticmethod
    def create_empty_excel():
        """Cria arquivo Excel vazio"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 'Empty'

        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)

        return SimpleUploadedFile(
            "test_empty.xlsx",
            excel_buffer.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    @staticmethod
    def get_standard_mapping():
        """Retorna mapeamento padrão para planilha complexa"""
        return {
            'code_column': 2,
            'description_column': 3,
            'dimensions_column': 4,
            'cubic_column': 5,
            'weight_column': 6,
            'ncm_column': 7,
            'data_start_row': 12,
            'price_columns': [
                {'index': 8, 'name': 'Preço 1'},
                {'index': 9, 'name': 'Preço 2'},
                {'index': 10, 'name': 'Preço 3'}
            ]
        }
