"""
Testes de integração para as views/endpoints do extractor
"""
from django.test import TestCase
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APIClient
from rest_framework import status
from extractor.models import SpreadsheetUpload, ColumnMapping, Product, ProductVariant
import io
import openpyxl
import json


class SpreadsheetUploadViewSetTest(TestCase):
    """Testes para SpreadsheetUploadViewSet"""

    def setUp(self):
        """Setup executado antes de cada teste"""
        self.client = APIClient()
        self.upload_url = '/api/uploads/'

        # Criar uma planilha Excel de teste real
        self.excel_file = self.create_test_excel()

    def create_test_excel(self):
        """Cria um arquivo Excel de teste válido"""
        wb = openpyxl.Workbook()
        ws = wb.active

        # Headers na linha 12 (index 11)
        headers = ['Foto', 'Num', 'Código', 'Descrição', 'Dimensões', 'Cúbico', 'Peso', 'NCM', 'Preço 1', 'Preço 2']
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=12, column=col_idx, value=header)

        # Dados
        data_rows = [
            [None, 1, 'AC2', 'Almofada Decorativa', 'A15 * L45 * P45', 0.037553, 0.55, '94049000', 123.12, 136.0],
            [None, 2, 'AC3', None, 'A15 * L55 * P55', 0.055233, 0.85, '94049000', 157.15, 173.0],
            [None, 3, 'AC33', None, 'A15 * L60 * P60', 0.066, 1.20, '94049000', 189.00, 205.0],
            [None, 4, 'BC1', 'Banqueta', 'A30 * L30 * P30', 0.027, 2.50, '94036000', 250.00, 275.0],
        ]

        for row_idx, row_data in enumerate(data_rows, start=13):
            for col_idx, value in enumerate(row_data, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        # Salvar em BytesIO
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)

        return SimpleUploadedFile(
            "test.xlsx",
            excel_buffer.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    def test_upload_spreadsheet_success(self):
        """Testa upload de planilha com sucesso"""
        response = self.client.post(
            self.upload_url,
            {'file': self.excel_file},
            format='multipart'
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertIn('id', response.data)
        self.assertIn('original_filename', response.data)
        self.assertEqual(response.data['original_filename'], 'test.xlsx')
        self.assertGreater(response.data['total_rows'], 0)
        self.assertGreater(response.data['total_columns'], 0)

    def test_upload_without_file(self):
        """Testa upload sem enviar arquivo"""
        response = self.client.post(self.upload_url, {}, format='multipart')

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('error', response.data)

    def test_list_uploads(self):
        """Testa listagem de uploads"""
        # Criar upload
        upload = SpreadsheetUpload.objects.create(
            file=self.excel_file,
            original_filename="test.xlsx",
            total_rows=10,
            total_columns=5
        )

        response = self.client.get(self.upload_url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertGreater(len(response.data), 0)

    def test_retrieve_upload(self):
        """Testa busca de um upload específico"""
        upload = SpreadsheetUpload.objects.create(
            file=self.excel_file,
            original_filename="test.xlsx",
            total_rows=10,
            total_columns=5
        )

        response = self.client.get(f'{self.upload_url}{upload.id}/')

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['id'], upload.id)
        self.assertEqual(response.data['original_filename'], 'test.xlsx')

    def test_preview_spreadsheet(self):
        """Testa preview da planilha"""
        # Criar arquivo Excel novamente para upload
        excel_file = self.create_test_excel()

        response = self.client.post(
            self.upload_url,
            {'file': excel_file},
            format='multipart'
        )
        upload_id = response.data['id']

        # Buscar preview
        preview_url = f'{self.upload_url}{upload_id}/preview/'
        response = self.client.get(preview_url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('total_rows', response.data)
        self.assertIn('total_columns', response.data)
        self.assertIn('preview', response.data)
        self.assertIsInstance(response.data['preview'], list)

    def test_process_spreadsheet_without_mapping(self):
        """Testa processamento sem mapeamento configurado"""
        excel_file = self.create_test_excel()

        response = self.client.post(
            self.upload_url,
            {'file': excel_file},
            format='multipart'
        )
        upload_id = response.data['id']

        # Tentar processar sem mapping
        process_url = f'{self.upload_url}{upload_id}/process/'
        response = self.client.post(process_url)

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('error', response.data)

    def test_process_spreadsheet_with_mapping(self):
        """Testa processamento completo com mapeamento"""
        # Upload
        excel_file = self.create_test_excel()
        response = self.client.post(
            self.upload_url,
            {'file': excel_file},
            format='multipart'
        )
        upload_id = response.data['id']

        # Criar mapeamento
        mapping_data = {
            'upload': upload_id,
            'code_column': 2,
            'description_column': 3,
            'dimensions_column': 4,
            'cubic_column': 5,
            'weight_column': 6,
            'ncm_column': 7,
            'data_start_row': 12,
            'price_columns': [
                {'index': 8, 'name': 'Preço 1'},
                {'index': 9, 'name': 'Preço 2'}
            ]
        }

        mapping_url = '/api/mappings/create_or_update/'
        response = self.client.post(
            mapping_url,
            data=json.dumps(mapping_data),
            content_type='application/json'
        )

        # Processar
        process_url = f'{self.upload_url}{upload_id}/process/'
        response = self.client.post(process_url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('products_created', response.data)
        self.assertIn('variants_created', response.data)
        self.assertGreater(response.data['products_created'], 0)
        self.assertGreater(response.data['variants_created'], 0)

    def test_export_json(self):
        """Testa exportação em formato JSON"""
        # Criar dados processados
        upload = SpreadsheetUpload.objects.create(
            file=self.excel_file,
            original_filename="test.xlsx"
        )
        product = Product.objects.create(
            upload=upload,
            description="Almofada"
        )
        ProductVariant.objects.create(
            product=product,
            upload=upload,
            code="AC2",
            dimensions="45x45",
            row_number=12,
            prices={'Preço 1': 123.12}
        )

        # Exportar
        export_url = f'{self.upload_url}{upload.id}/export/?format=json'
        response = self.client.get(export_url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIsInstance(response.data, list)
        self.assertGreater(len(response.data), 0)

    def test_export_csv(self):
        """Testa exportação em formato CSV"""
        upload = SpreadsheetUpload.objects.create(
            file=self.excel_file,
            original_filename="test.xlsx"
        )
        product = Product.objects.create(
            upload=upload,
            description="Almofada"
        )
        ProductVariant.objects.create(
            product=product,
            upload=upload,
            code="AC2",
            dimensions="45x45",
            row_number=12
        )

        export_url = f'{self.upload_url}{upload.id}/export/?format=csv'
        response = self.client.get(export_url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response['Content-Type'], 'text/csv; charset=utf-8')

    def test_export_xml(self):
        """Testa exportação em formato XML"""
        upload = SpreadsheetUpload.objects.create(
            file=self.excel_file,
            original_filename="test.xlsx"
        )
        product = Product.objects.create(
            upload=upload,
            description="Almofada"
        )
        ProductVariant.objects.create(
            product=product,
            upload=upload,
            code="AC2",
            dimensions="45x45",
            row_number=12
        )

        export_url = f'{self.upload_url}{upload.id}/export/?format=xml'
        response = self.client.get(export_url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response['Content-Type'], 'application/xml')

    def test_export_invalid_format(self):
        """Testa exportação com formato inválido"""
        upload = SpreadsheetUpload.objects.create(
            file=self.excel_file,
            original_filename="test.xlsx"
        )

        export_url = f'{self.upload_url}{upload.id}/export/?format=invalid'
        response = self.client.get(export_url)

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)


class ColumnMappingViewSetTest(TestCase):
    """Testes para ColumnMappingViewSet"""

    def setUp(self):
        """Setup executado antes de cada teste"""
        self.client = APIClient()
        self.mapping_url = '/api/mappings/'

        excel_file = SimpleUploadedFile("test.xlsx", b"fake content")
        self.upload = SpreadsheetUpload.objects.create(
            file=excel_file,
            original_filename="test.xlsx"
        )

    def test_create_mapping(self):
        """Testa criação de mapeamento"""
        data = {
            'upload': self.upload.id,
            'code_column': 2,
            'description_column': 3,
            'data_start_row': 12
        }

        response = self.client.post(
            self.mapping_url,
            data=json.dumps(data),
            content_type='application/json'
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['code_column'], 2)

    def test_create_or_update_new_mapping(self):
        """Testa criação de novo mapeamento via create_or_update"""
        data = {
            'upload': self.upload.id,
            'code_column': 2,
            'description_column': 3,
            'dimensions_column': 4,
            'data_start_row': 12,
            'price_columns': [
                {'index': 8, 'name': 'Preço Fornecido'}
            ]
        }

        url = f'{self.mapping_url}create_or_update/'
        response = self.client.post(
            url,
            data=json.dumps(data),
            content_type='application/json'
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['code_column'], 2)
        self.assertEqual(len(response.data['price_columns']), 1)

    def test_create_or_update_existing_mapping(self):
        """Testa atualização de mapeamento existente"""
        # Criar mapeamento inicial
        ColumnMapping.objects.create(
            upload=self.upload,
            code_column=2,
            data_start_row=10
        )

        # Atualizar
        data = {
            'upload': self.upload.id,
            'code_column': 3,  # Mudou
            'data_start_row': 15  # Mudou
        }

        url = f'{self.mapping_url}create_or_update/'
        response = self.client.post(
            url,
            data=json.dumps(data),
            content_type='application/json'
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['code_column'], 3)
        self.assertEqual(response.data['data_start_row'], 15)

        # Verificar que só existe um mapping
        self.assertEqual(ColumnMapping.objects.filter(upload=self.upload).count(), 1)

    def test_create_mapping_without_upload(self):
        """Testa criação sem especificar upload"""
        data = {
            'code_column': 2
        }

        url = f'{self.mapping_url}create_or_update/'
        response = self.client.post(
            url,
            data=json.dumps(data),
            content_type='application/json'
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_list_mappings(self):
        """Testa listagem de mapeamentos"""
        ColumnMapping.objects.create(
            upload=self.upload,
            code_column=2
        )

        response = self.client.get(self.mapping_url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertGreater(len(response.data), 0)


class ProductViewSetTest(TestCase):
    """Testes para ProductViewSet"""

    def setUp(self):
        """Setup executado antes de cada teste"""
        self.client = APIClient()
        self.product_url = '/api/products/'

        excel_file = SimpleUploadedFile("test.xlsx", b"fake content")
        self.upload = SpreadsheetUpload.objects.create(
            file=excel_file,
            original_filename="test.xlsx"
        )

        self.product = Product.objects.create(
            upload=self.upload,
            description="Almofada Decorativa"
        )

        ProductVariant.objects.create(
            product=self.product,
            upload=self.upload,
            code="AC2",
            dimensions="45x45",
            row_number=12
        )
        ProductVariant.objects.create(
            product=self.product,
            upload=self.upload,
            code="AC3",
            dimensions="55x55",
            row_number=13
        )

    def test_list_products(self):
        """Testa listagem de produtos"""
        response = self.client.get(self.product_url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertGreater(len(response.data), 0)

    def test_list_products_by_upload(self):
        """Testa listagem filtrada por upload_id"""
        # Criar outro upload com produtos
        other_upload = SpreadsheetUpload.objects.create(
            file=SimpleUploadedFile("other.xlsx", b"content"),
            original_filename="other.xlsx"
        )
        Product.objects.create(
            upload=other_upload,
            description="Outro Produto"
        )

        # Filtrar por upload
        url = f'{self.product_url}?upload_id={self.upload.id}'
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        # Response pode estar paginado
        if isinstance(response.data, dict) and 'results' in response.data:
            self.assertEqual(len(response.data['results']), 1)
            self.assertEqual(response.data['results'][0]['description'], "Almofada Decorativa")
        else:
            self.assertEqual(len(response.data), 1)
            self.assertEqual(response.data[0]['description'], "Almofada Decorativa")

    def test_retrieve_product(self):
        """Testa busca de produto específico"""
        url = f'{self.product_url}{self.product.id}/'
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['id'], self.product.id)
        self.assertEqual(response.data['description'], "Almofada Decorativa")

    def test_product_includes_variants(self):
        """Testa se produto inclui variantes"""
        url = f'{self.product_url}{self.product.id}/'
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('variants', response.data)
        self.assertEqual(len(response.data['variants']), 2)

    def test_product_variants_data(self):
        """Testa dados das variantes"""
        url = f'{self.product_url}{self.product.id}/'
        response = self.client.get(url)

        variants = response.data['variants']
        codes = [v['code'] for v in variants]

        self.assertIn('AC2', codes)
        self.assertIn('AC3', codes)


class FullWorkflowIntegrationTest(TestCase):
    """Testa o fluxo completo de upload até exportação"""

    def setUp(self):
        """Setup executado antes de cada teste"""
        self.client = APIClient()

    def create_test_excel(self):
        """Cria planilha Excel de teste"""
        wb = openpyxl.Workbook()
        ws = wb.active

        headers = ['Foto', 'Num', 'Código', 'Descrição', 'Dimensões', 'Cúbico', 'Peso', 'NCM', 'Preço']
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=12, column=col_idx, value=header)

        data_rows = [
            [None, 1, 'AC2', 'Almofada', '45x45', 0.037, 0.55, '94049000', 123.12],
            [None, 2, 'AC3', None, '55x55', 0.055, 0.85, '94049000', 157.15],
        ]

        for row_idx, row_data in enumerate(data_rows, start=13):
            for col_idx, value in enumerate(row_data, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)

        return SimpleUploadedFile("test.xlsx", excel_buffer.read())

    def test_complete_workflow(self):
        """Testa fluxo completo: upload -> mapping -> process -> export"""
        # 1. Upload
        upload_response = self.client.post(
            '/api/uploads/',
            {'file': self.create_test_excel()},
            format='multipart'
        )
        self.assertEqual(upload_response.status_code, status.HTTP_201_CREATED)
        upload_id = upload_response.data['id']

        # 2. Preview
        preview_response = self.client.get(f'/api/uploads/{upload_id}/preview/')
        self.assertEqual(preview_response.status_code, status.HTTP_200_OK)

        # 3. Configurar mapeamento
        mapping_data = {
            'upload': upload_id,
            'code_column': 2,
            'description_column': 3,
            'dimensions_column': 4,
            'cubic_column': 5,
            'weight_column': 6,
            'ncm_column': 7,
            'data_start_row': 12,
            'price_columns': [{'index': 8, 'name': 'Preço'}]
        }
        mapping_response = self.client.post(
            '/api/mappings/create_or_update/',
            data=json.dumps(mapping_data),
            content_type='application/json'
        )
        self.assertEqual(mapping_response.status_code, status.HTTP_200_OK)

        # 4. Processar
        process_response = self.client.post(f'/api/uploads/{upload_id}/process/')
        self.assertEqual(process_response.status_code, status.HTTP_200_OK)
        self.assertEqual(process_response.data['products_created'], 1)
        self.assertEqual(process_response.data['variants_created'], 2)

        # 5. Listar produtos
        products_response = self.client.get(f'/api/products/?upload_id={upload_id}')
        self.assertEqual(products_response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(products_response.data), 1)
        self.assertEqual(len(products_response.data[0]['variants']), 2)

        # 6. Exportar JSON
        export_json = self.client.get(f'/api/uploads/{upload_id}/export/?format=json')
        self.assertEqual(export_json.status_code, status.HTTP_200_OK)
        self.assertEqual(len(export_json.data), 2)

        # 7. Exportar CSV
        export_csv = self.client.get(f'/api/uploads/{upload_id}/export/?format=csv')
        self.assertEqual(export_csv.status_code, status.HTTP_200_OK)

        # 8. Exportar XML
        export_xml = self.client.get(f'/api/uploads/{upload_id}/export/?format=xml')
        self.assertEqual(export_xml.status_code, status.HTTP_200_OK)
